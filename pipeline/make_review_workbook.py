"""Build the analyst REVIEW WORKBOOK for one client - the agreement check, and the manual
preview of the PIDA review screen.

Design settled with Sameer 2026-08-06 (RUN_LOG, and DEPLOYMENT-CONCEPT section 3):

  * `qa_line` keeps its 40+ columns. The ANALYST sees a reduced surface - his instruction:
    "the reduced cols when deployed to the app should come from a view with limited cols so our
    analyst dont get overwhelmed with reading 40+ cols". This file writes exactly that surface.
  * MACHINERY IS HIDDEN, NOT ABSENT. `qa_line_id` travels in a hidden column so answers can be
    re-attached after the reviewer sorts, filters or edits. `rules_table` is NOT in the file at
    all - the analyst decides the category, they never route a fix; routing is the proposal
    writer's job. Sameer: "this would be useful for you but not the analyst."
  * `gl_account_name` IS on the surface, and it earned its place by measurement: 161 of the
    2,000 judged pilot lines carry a Correct/Incorrect verdict on a line with NO usable item
    description (Western 92, Sydney Adventist 58, Melbourne 11). Without the GL column a
    reviewer has nothing to check those verdicts against.
  * THE CATEGORY PICKER IS RESTRICTED BY THE FILE ITSELF. The dropdown is fed from a hidden
    sheet holding ONLY this hospital's taxonomy, so choosing another client's category is not
    possible - the isolation rule enforced by data validation rather than by trust. And the list
    differs per row: in-scope indirect categories on normal lines, the FULL taxonomy on
    uncategorised lines, which is the per-line scope guard `judge.apply_verdicts` already
    applies.

Known Excel limit, stated rather than discovered: a cell carries ONE validation rule, so the
category cell cannot be hard-locked until "Disagree" is chosen. It is highlighted and prompted,
not enforced.

Usage:
    python pipeline/make_review_workbook.py --client sydney_adventist
    python pipeline/make_review_workbook.py --client sydney_adventist --outdir QA_LINE_TEST
"""
import argparse
import datetime
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

import clientcfg
from db import connect_qa, load_env

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# --- the analyst surface, in reading order. (header, source column, width) ---------------------
# qa_line_id leads and is HIDDEN - it is the round-trip key, not information for the reviewer.
COLUMNS = [
    ("qa_line_id",              "QA_LINE_ID",                 10),
    ("Client",                  "CLIENT_CODE",                18),
    ("Supplier",                "SUPPLIER_NAME",              34),
    ("Item description",        "ITEM_DESCRIPTION",           44),
    ("GL account name",         "GL_ACCOUNT_NAME",            26),
    ("Spend",                   "SPEND",                      14),
    ("Current L0",              "CATEGORY_LVL_0",             16),
    ("Current L1",              "CATEGORY_LVL_1",             24),
    ("Current L2",              "CATEGORY_LVL_2",             26),
    ("Current L3",              "CATEGORY_LVL_3",             26),
    ("Current L4",              "CATEGORY_LVL_4",             26),
    ("Rule ID",                 "RULE_ID",                    14),
    ("Rule priority",           "RULE_PRIORITY",              13),
    ("Our verdict",             "VERDICT",                    12),
    ("Confidence",              "CONFIDENCE",                 11),
    ("Suggested L0",            "SUGGESTED_CATEGORY_LVL_0",   16),
    ("Suggested L1",            "SUGGESTED_CATEGORY_LVL_1",   24),
    ("Suggested L2",            "SUGGESTED_CATEGORY_LVL_2",   26),
    ("Suggested L3",            "SUGGESTED_CATEGORY_LVL_3",   26),
    ("Suggested L4",            "SUGGESTED_CATEGORY_LVL_4",   26),
    ("Basis",                   "BASIS",                      22),
    ("Why we said that",        "RATIONALE",                  70),
    # --- the reviewer's columns -----------------------------------------------------------
    ("YOUR RESPONSE",           None,                         34),
    ("CORRECT CATEGORY (if disagreeing)", None,               60),
    ("YOUR NOTE (optional)",    None,                         44),
]

HIDDEN_COL = 1                       # qa_line_id
COL_RESPONSE = len(COLUMNS) - 2      # 1-based positions computed below
COL_CATEGORY = len(COLUMNS) - 1
COL_NOTE = len(COLUMNS)

RESPONSE_OPTIONS = [
    "Agree",
    "Disagree - current category is fine",
    "Disagree - correct category is:",
]

# Plain hyphens, not dashes: the labels are embedded in a validation formula and a conditional
# formatting formula, and punctuation that round-trips badly breaks both silently.
DISAGREE_WITH_CATEGORY = RESPONSE_OPTIONS[2]

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
ASK_FILL = PatternFill("solid", fgColor="C00000")
PROMPT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

VERDICT_ORDER = {"Incorrect": 0, "Uncertain": 1, "Correct": 2}


def fetch(client):
    """Read the judged lines and this client's taxonomy. One connection, pilot database only."""
    qa = connect_qa()
    cur = qa.cursor()

    cur.execute("SELECT DISTINCT run_id FROM qa_line")
    runs = [r[0] for r in cur.fetchall()]
    if len(runs) != 1:
        raise SystemExit(
            f"STOP: expected exactly one run_id in qa_line, found {len(runs)}: {runs}. "
            "A superseded generation is still in the table - see CLAUDE.md."
        )
    run_id = runs[0]

    select = ", ".join(f"[{src}]" for _, src, _ in COLUMNS if src)
    cur.execute(
        f"""SELECT {select},
                   CASE WHEN LTRIM(RTRIM(ISNULL(CATEGORY_LVL_1,''))) = ? THEN 1 ELSE 0 END
                     AS is_uncategorised
              FROM qa_line WHERE client_code = ?""",
        clientcfg.LEVEL_UNCATEGORISED, client)
    rows = [list(r) for r in cur.fetchall()]
    if not rows:
        raise SystemExit(f"STOP: no rows in qa_line for client {client!r}.")

    # Incorrect first (the consequential calls), then Uncertain, then Correct; supplier within.
    vi = [i for i, (_, src, _) in enumerate(COLUMNS) if src == "VERDICT"][0]
    si = [i for i, (_, src, _) in enumerate(COLUMNS) if src == "SUPPLIER_NAME"][0]
    di = [i for i, (_, src, _) in enumerate(COLUMNS) if src == "ITEM_DESCRIPTION"][0]
    rows.sort(key=lambda r: (VERDICT_ORDER.get(r[vi], 9), (r[si] or ""), (r[di] or "")))

    cur.execute("""SELECT DISTINCT path_full, in_scope FROM qa_category
                    WHERE client_code = ? AND path_full IS NOT NULL
                    ORDER BY path_full""", client)
    tax = cur.fetchall()
    qa.close()

    in_scope = [t[0] for t in tax if t[1]]
    all_paths = [t[0] for t in tax]
    if not in_scope:
        raise SystemExit(f"STOP: no in-scope categories loaded for {client!r}.")
    return run_id, rows, in_scope, all_paths


def ranges_for(rows_1based, col_letter):
    """Compress row numbers into an Excel sqref like 'X2:X40 X55:X60'."""
    if not rows_1based:
        return ""
    out, start, prev = [], rows_1based[0], rows_1based[0]
    for r in rows_1based[1:]:
        if r == prev + 1:
            prev = r
            continue
        out.append((start, prev))
        start = prev = r
    out.append((start, prev))
    return " ".join(f"{col_letter}{a}:{col_letter}{b}" if a != b else f"{col_letter}{a}"
                    for a, b in out)


def build(client, outdir_name):
    load_env()
    run_id, rows, in_scope, all_paths = fetch(client)
    display = client.replace("_", " ").title()
    today = datetime.date.today().isoformat()

    wb = Workbook()

    # ---------------------------------------------------------------- sheet 1: Read me
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 108
    readme = [
        (f"Indirect Spend QA - Review Workbook - {display}", True),
        ("", False),
        (f"Generated {today} from run {run_id}. {len(rows):,} judged lines.", False),
        ("", False),
        ("WHAT THIS IS", True),
        ("Every line here was judged automatically against your hospital's own category "
         "structure. This file is how a person checks that work before any accuracy figure is "
         "quoted to a client.", False),
        ("", False),
        ("HOW TO REVIEW", True),
        ("1. Go to the 'Review' sheet. Lines we marked Incorrect come first, then Uncertain, "
         "then Correct.", False),
        ("2. Read the line (supplier, item description, GL account), the category it currently "
         "sits in, and what we said about it.", False),
        ("3. In 'YOUR RESPONSE', choose one of three:", False),
        ("      Agree  -  our verdict is right.", False),
        ("      Disagree - current category is fine  -  we called it wrong, but it was not.", False),
        (f"      {DISAGREE_WITH_CATEGORY}  -  then pick the right category in the next column.",
         False),
        ("4. Add a note if you want to explain. Free text, optional - but it is the most useful "
         "thing you can give us.", False),
        ("", False),
        ("ABOUT THE CATEGORY LIST", True),
        (f"The dropdown holds {display}'s categories and nothing else - no other hospital's "
         "categories can be chosen from this file. On normal lines it offers the indirect "
         "(non-clinical) categories. On lines that were never categorised at all it offers the "
         "full taxonomy, because those lines have not yet been through the clinical split.", False),
        ("", False),
        ("PLEASE DO NOT", True),
        ("Delete or reorder columns, or delete rows. Sorting and filtering are fine - each row "
         "carries a hidden reference that keeps your answers attached to the right line.", False),
        ("", False),
        ("WHAT HAPPENS NEXT", True),
        ("Your answers come back into the review database beside our findings - never on top of "
         "them. How often you disagreed is how we measure the judge. Where a correction is "
         "agreed, it becomes a proposed rule change for your hospital's own rules.", False),
    ]
    for i, (text, bold) in enumerate(readme, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            c.font = Font(bold=True, size=12 if i == 1 else 10)

    # ---------------------------------------------------------------- sheet 3: taxonomy (hidden)
    tx = wb.create_sheet("Taxonomy")
    tx["A1"] = "In-scope (indirect) categories"
    tx["B1"] = "All categories (uncategorised lines only)"
    for i, p in enumerate(in_scope, start=2):
        tx.cell(row=i, column=1, value=p)
    for i, p in enumerate(all_paths, start=2):
        tx.cell(row=i, column=2, value=p)
    tx.column_dimensions["A"].width = 90
    tx.column_dimensions["B"].width = 90
    tx.sheet_state = "hidden"

    # Defined names are the robust way to reference another sheet from data validation.
    wb.defined_names.add(DefinedName(
        "TaxInScope", attr_text=f"Taxonomy!$A$2:$A${len(in_scope) + 1}"))
    wb.defined_names.add(DefinedName(
        "TaxAll", attr_text=f"Taxonomy!$B$2:$B${len(all_paths) + 1}"))

    # ---------------------------------------------------------------- sheet 2: Review
    rv = wb.create_sheet("Review", 1)
    for ci, (header, _, width) in enumerate(COLUMNS, start=1):
        c = rv.cell(row=1, column=ci, value=header)
        c.fill = ASK_FILL if ci >= COL_RESPONSE else HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        c.border = BORDER
        rv.column_dimensions[get_column_letter(ci)].width = width
    rv.row_dimensions[1].height = 30

    src_count = sum(1 for _, s, _ in COLUMNS if s)
    uncat_rows, normal_rows = [], []
    for ri, row in enumerate(rows, start=2):
        is_uncat = bool(row[src_count])
        (uncat_rows if is_uncat else normal_rows).append(ri)
        vi = 0
        for ci, (_, src, _) in enumerate(COLUMNS, start=1):
            if src is None:
                c = rv.cell(row=ri, column=ci, value=None)
                c.fill = PROMPT_FILL
            else:
                c = rv.cell(row=ri, column=ci, value=row[vi])
                vi += 1
                if src == "SPEND":
                    c.number_format = '#,##0.00;[Red]-#,##0.00'
                elif src == "CONFIDENCE":
                    c.number_format = '0.00'
            c.alignment = Alignment(vertical="top", wrap_text=(ci in (4, 22, COL_NOTE)))
            c.border = BORDER

    last = len(rows) + 1
    rv.freeze_panes = "E2"
    rv.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"
    rv.column_dimensions[get_column_letter(HIDDEN_COL)].hidden = True

    resp_letter = get_column_letter(COL_RESPONSE)
    cat_letter = get_column_letter(COL_CATEGORY)

    from openpyxl.worksheet.datavalidation import DataValidation
    dv_resp = DataValidation(
        type="list", allow_blank=True,
        formula1='"' + ",".join(RESPONSE_OPTIONS) + '"',
        showDropDown=False, showErrorMessage=True,
        errorTitle="Pick from the list",
        error="Choose one of the three responses.",
        promptTitle="Your response",
        prompt="Agree, or say which kind of disagreement it is.")
    dv_resp.sqref = f"{resp_letter}2:{resp_letter}{last}"
    rv.add_data_validation(dv_resp)

    dv_scope = DataValidation(
        type="list", allow_blank=True, formula1="=TaxInScope",
        showDropDown=False, showErrorMessage=True,
        errorTitle="Not a category of this hospital",
        error="Pick from the dropdown. It holds this hospital's indirect categories only.",
        promptTitle="Correct category",
        prompt="Only needed if you chose 'Disagree - correct category is:'.")
    if normal_rows:
        dv_scope.sqref = ranges_for(normal_rows, cat_letter)
        rv.add_data_validation(dv_scope)

    dv_all = DataValidation(
        type="list", allow_blank=True, formula1="=TaxAll",
        showDropDown=False, showErrorMessage=True,
        errorTitle="Not a category of this hospital",
        error="Pick from the dropdown. This line was never categorised, so the full taxonomy "
              "is offered - clinical branches included.",
        promptTitle="Correct category (full taxonomy)",
        prompt="This line has no category at all, so any branch may be chosen.")
    if uncat_rows:
        dv_all.sqref = ranges_for(uncat_rows, cat_letter)
        rv.add_data_validation(dv_all)

    # Prompt the category cell when - and only when - a category is actually being asked for.
    rv.conditional_formatting.add(
        f"{cat_letter}2:{cat_letter}{last}",
        FormulaRule(formula=[f'${resp_letter}2="{DISAGREE_WITH_CATEGORY}"'],
                    fill=PatternFill("solid", fgColor="FFE699"), stopIfTrue=False))

    # ---------------------------------------------------------------- write
    outdir = os.path.join(PROJECT_ROOT, "output", outdir_name)
    os.makedirs(outdir, exist_ok=True)
    # "Indirect" appears in every artefact name - the folder is Medical QA - Indirects, and a
    # reader who loses that word concludes we QA'd medical items, which is the opposite.
    path = os.path.join(outdir, f"Indirect QA Review - {display} - {today}.xlsx")
    wb.save(path)

    print(f"  client        : {client}  ({display})")
    print(f"  run_id        : {run_id}")
    print(f"  lines written : {len(rows):,}   uncategorised (full-taxonomy picker): "
          f"{len(uncat_rows)}")
    print(f"  dropdown sizes: in-scope {len(in_scope):,} | full taxonomy {len(all_paths):,}")
    print(f"  columns       : {len(COLUMNS)} ({len(COLUMNS) - 1} visible, qa_line_id hidden)")
    print(f"  written to    : {path}")
    return path


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Build the analyst review workbook for one client.")
    ap.add_argument("--client", required=True, help="client key, e.g. the folder under clients/")
    ap.add_argument("--outdir", default=None,
                    help="folder name under output/ (default: the client key)")
    a = ap.parse_args()
    build(a.client, a.outdir or a.client)
