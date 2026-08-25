"""PHASE 0 - profile every hospital client and recommend the POC.

Run:  python pipeline/profile_clients.py            # all clients
      python pipeline/profile_clients.py western_health

Answers the questions the plan says must be answered before any client run:
  - Do the credentials work for all four?
  - How many lines / how much spend / how many vendors each?
  - What shape is each hospital's taxonomy, and which categories are clinical?
  - How good are the item descriptions?  (the judge reads these - poor descriptions cap the
    achievable accuracy, so this is the main POC selection criterion)
  - How many distinct judge units - i.e. what does a census actually cost?
  - Is this client in the MSD, and what's its coherence split?

Writes program/Client Profiling.xlsx. Nothing here mutates client data - reads only.

Where a client's config.yaml has no source.table yet, the profiler lists candidate tables and
views with row counts instead, so the table can be chosen from evidence.
"""
import os
import sys
import traceback
from datetime import datetime

import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import clientcfg  # noqa: E402
import msd as msdmod  # noqa: E402
from db import connect_client, connect_msd, load_env  # noqa: E402

PROJECT_ROOT = clientcfg.PROJECT_ROOT
OUT_PATH = os.path.join(PROJECT_ROOT, "program", "Client Profiling.xlsx")

# Descriptions that carry no categorisation signal. A high share of these means the judge is
# reasoning from the vendor alone, and the honest verdict for those units is Uncertain.
GENERIC_TERMS = {
    "GOODS", "SERVICES", "SERVICE", "MISC", "MISCELLANEOUS", "SUNDRY", "SUNDRIES",
    "PURCHASE", "INVOICE", "CHARGES", "CHARGE", "N/A", "NA", "NONE", "OTHER",
    "GENERAL", "ITEMS", "ITEM", "SUPPLY", "SUPPLIES", "PRODUCT", "PRODUCTS",
    "CONSULTING", "MONTHLY CHARGE", "ADJUSTMENT", "CREDIT", "FREIGHT",
}


def q(col):
    """Quote a SQL identifier."""
    return "[" + str(col).replace("]", "]]") + "]"


def _clean(col):
    """Trimmed, NULL-if-blank text form of one description column."""
    return f"NULLIF(LTRIM(RTRIM(CONVERT(nvarchar(900), {q(col)}))), '')"


def desc_expr(fields):
    """The composed line description: first non-blank of the client's description_fields.

    Clients differ - Sydney Adventist has no ITEM_DESCRIPTION at all and uses INVOICE
    DESCRIPTION; Melbourne has PO and invoice text that fill gaps where ITEM_DESCRIPTION is
    blank. So description is a priority LIST, not a column, and this is what the judge reads.
    """
    if not fields:
        return "CAST(NULL AS nvarchar(900))"
    if len(fields) == 1:
        return _clean(fields[0])
    return "COALESCE(" + ", ".join(_clean(f) for f in fields) + ")"


def desc_source_expr(fields):
    """Which of the description_fields actually supplied the text - becomes desc_source_field."""
    if not fields:
        return "'(none)'"
    whens = " ".join(
        "WHEN {} IS NOT NULL THEN '{}'".format(_clean(f), str(f)[:60].replace("'", "''"))
        for f in fields)
    return f"CASE {whens} ELSE '(all blank)' END"


def scalar_df(cn, sql):
    return pd.read_sql(sql, cn)


# ----------------------------------------------------------------------------- discovery ----
def discover_tables(cn):
    """Candidate fact tables/views with row counts - used when source.table isn't set yet."""
    sql = """
    SELECT TOP 60
           s.name AS [schema], t.name AS [table], 'TABLE' AS kind,
           SUM(p.rows) AS approx_rows
    FROM sys.tables t
    JOIN sys.schemas s ON s.schema_id = t.schema_id
    JOIN sys.partitions p ON p.object_id = t.object_id AND p.index_id IN (0,1)
    GROUP BY s.name, t.name
    UNION ALL
    SELECT TOP 60 s.name, v.name, 'VIEW', NULL
    FROM sys.views v JOIN sys.schemas s ON s.schema_id = v.schema_id
    ORDER BY approx_rows DESC
    """
    return pd.read_sql(sql, cn)


def table_columns(cn, table):
    """Column names + types for a configured table, so the column map can be checked."""
    bare = table.replace("[", "").replace("]", "")
    schema, _, name = bare.rpartition(".")
    schema = schema or "dbo"
    sql = """
        SELECT COLUMN_NAME, DATA_TYPE
        FROM INFORMATION_SCHEMA.COLUMNS
        WHERE TABLE_SCHEMA = ? AND TABLE_NAME = ?
        ORDER BY ORDINAL_POSITION
    """
    return pd.read_sql(sql, cn, params=[schema, name])


# ------------------------------------------------------------------------------ profiling ----
def profile_volume(cn, table, cols, d):
    """Volume, spend, vendors, credits AND description quality - deliberately one scan.

    These were two separate full scans. On a 3.4M-row view that is minutes of duplicated work,
    so they are merged: every aggregate here is computed in a single pass.
    """
    sup, spend = q(cols["supplier"]), q(cols["spend"])
    cat0 = q(cols["cat_l0"])
    sql = f"""
        SELECT COUNT(*) AS lines,
               SUM(TRY_CONVERT(float, {spend})) AS spend,
               SUM(ABS(TRY_CONVERT(float, {spend}))) AS spend_abs,
               COUNT(DISTINCT UPPER(LTRIM(RTRIM(CONVERT(nvarchar(400), {sup}))))) AS vendors,
               SUM(CASE WHEN TRY_CONVERT(float, {spend}) IS NULL THEN 1 ELSE 0 END) AS spend_unparseable,
               SUM(CASE WHEN TRY_CONVERT(float, {spend}) < 0 THEN 1 ELSE 0 END) AS spend_negative,
               SUM(CASE WHEN TRY_CONVERT(float, {spend}) < 0
                        THEN ABS(TRY_CONVERT(float, {spend})) ELSE 0 END) AS credit_spend_abs,
               SUM(CASE WHEN {cat0} IS NULL THEN 1 ELSE 0 END) AS cat0_null,
               -- description quality, same pass
               SUM(CASE WHEN {d} IS NULL THEN 1 ELSE 0 END) AS blank,
               SUM(CASE WHEN LEN({d}) BETWEEN 1 AND 7 THEN 1 ELSE 0 END) AS very_short,
               SUM(CASE WHEN {d} IS NOT NULL
                         AND {d} NOT LIKE '%[^0-9 .,/-]%' THEN 1 ELSE 0 END) AS numeric_only,
               SUM(CASE WHEN {d} LIKE '%[0-9][0-9][0-9][0-9][0-9][0-9]%'
                        THEN 1 ELSE 0 END) AS long_digit_run,
               SUM(CASE WHEN {d} LIKE '%MRN%' OR {d} LIKE '%UR NO%' OR {d} LIKE '%U/R%'
                          OR {d} LIKE '%D.O.B%' OR {d} LIKE '%DOB %' THEN 1 ELSE 0 END) AS pii_marker
        FROM {table}
    """
    return scalar_df(cn, sql).iloc[0].to_dict()


def profile_units(cn, table, cols, d):
    """Distinct judge units at (vendor, term, assigned category) - the census cost driver.

    Both grains come from ONE grouped pass. Counting the (vendor, term) grain separately meant a
    second full scan for a number derivable from the first: group the fine grain, then count
    distinct pairs within it.

    The gap between the two is the Cleanaway Case B population - same vendor, same term, sitting
    in two different categories. That gap is a rule-application defect that a (vendor, term) unit
    would collapse into a single verdict and hide.
    """
    sup, cat = q(cols["supplier"]), q(cols["cat_l2"])
    v = f"UPPER(LTRIM(RTRIM(CONVERT(nvarchar(400), {sup}))))"
    sql = f"""
        SELECT COUNT(*) AS units,
               COUNT(DISTINCT ISNULL(v,'') + NCHAR(31) + ISNULL(t,'')) AS units_vt
        FROM (
            SELECT {v} AS v, {d} AS t, CONVERT(nvarchar(400), {cat}) AS c
            FROM {table}
            GROUP BY {v}, {d}, CONVERT(nvarchar(400), {cat})
        ) x
    """
    r = scalar_df(cn, sql).iloc[0]
    return int(r["units"]), int(r["units_vt"])


def profile_desc_fields(cn, table, fields):
    """Fill rate of EACH candidate description column, and which one actually wins.

    The plan asks for description quality 'per candidate field'. A composed figure hides the
    answer that matters: whether the fallback columns are earning their place, or whether one
    field is carrying everything.
    """
    if not fields:
        return pd.DataFrame()
    sel = ",\n               ".join(
        f"SUM(CASE WHEN {_clean(f)} IS NOT NULL THEN 1 ELSE 0 END) AS {q('f%d' % i)}"
        for i, f in enumerate(fields))
    src = desc_source_expr(fields)
    sql = f"""
        SELECT COUNT(*) AS lines,
               {sel},
               SUM(CASE WHEN {src} = '(all blank)' THEN 1 ELSE 0 END) AS all_blank
        FROM {table}
    """
    r = scalar_df(cn, sql).iloc[0]
    n = max(int(r["lines"]), 1)
    rows = [{"field": f, "populated_lines": int(r[f"f{i}"]),
             "populated_pct": round(100.0 * int(r[f"f{i}"]) / n, 1)}
            for i, f in enumerate(fields)]
    # Label must match desc_source_expr's ELSE branch exactly - the merge below joins on it.
    rows.append({"field": "(all blank)", "populated_lines": int(r["all_blank"]),
                 "populated_pct": round(100.0 * int(r["all_blank"]) / n, 1)})

    # which field actually supplies the text, after priority is applied
    sql2 = f"""
        SELECT {src} AS won_by, COUNT(*) AS lines
        FROM {table} GROUP BY {src}
    """
    won = pd.read_sql(sql2, cn)
    out = pd.DataFrame(rows)
    out = out.merge(won.rename(columns={"won_by": "field", "lines": "supplied_the_text"}),
                    on="field", how="left")
    out["supplied_the_text"] = out["supplied_the_text"].fillna(0).astype(int)
    return out


def top_descriptions(cn, table, cols, d, n=300):
    """The composed descriptions carrying the most spend - what the judge will actually read.

    Ordered by ABS(spend): credits are real lines that need judging, and ordering by signed spend
    would push a large credit to the bottom of the list instead of surfacing it.
    """
    spend = q(cols["spend"])
    sql = f"""
        SELECT TOP {n}
               UPPER({d}) AS item_desc,
               COUNT(*) AS lines,
               SUM(TRY_CONVERT(float, {spend})) AS spend
        FROM {table}
        GROUP BY UPPER({d})
        ORDER BY SUM(ABS(TRY_CONVERT(float, {spend}))) DESC
    """
    return pd.read_sql(sql, cn)


def taxonomy_by_level(cn, table, cols):
    """Distinct value count at each taxonomy level the client has mapped."""
    rows = []
    for lvl in ["cat_l0", "cat_l1", "cat_l2", "cat_l3", "cat_l4"]:
        if lvl not in cols:
            continue
        c = q(cols[lvl])
        sql = f"""
            SELECT COUNT(DISTINCT CONVERT(nvarchar(400), {c})) AS distinct_values,
                   SUM(CASE WHEN {c} IS NULL OR LTRIM(RTRIM(CONVERT(nvarchar(400), {c}))) = ''
                            THEN 1 ELSE 0 END) AS blank_lines
            FROM {table}
        """
        r = scalar_df(cn, sql).iloc[0]
        rows.append({"level": lvl, "client_column": cols[lvl],
                     "distinct_values": int(r["distinct_values"]),
                     "blank_lines": int(r["blank_lines"])})
    return pd.DataFrame(rows)


def taxonomy_detail(cn, table, cols):
    """Every category value with lines and spend. This is what you read to decide which
    categories are clinical (the clinical gate is category-based only, by design)."""
    spend = q(cols["spend"])
    frames = []
    for lvl in ["cat_l1", "cat_l2", "cat_l3"]:
        if lvl not in cols:
            continue
        c = q(cols[lvl])
        sql = f"""
            SELECT '{lvl}' AS level,
                   CONVERT(nvarchar(400), {c}) AS category,
                   COUNT(*) AS lines,
                   SUM(TRY_CONVERT(float, {spend})) AS spend
            FROM {table}
            GROUP BY CONVERT(nvarchar(400), {c})
        """
        frames.append(pd.read_sql(sql, cn))
    if not frames:
        return pd.DataFrame()
    out = pd.concat(frames, ignore_index=True)
    return out.sort_values(["level", "spend"], ascending=[True, False])


def profile_msd(cfg, env):
    """MSD status for this client: is it enriched, and what's the coherence split?"""
    code = clientcfg.msd_client_code(cfg)
    result = {"msd_client_code": code or "(not set)", "in_msd": False,
              "msd_vendors": 0, "coherence": pd.DataFrame()}
    if not code:
        return result
    cid = msdmod.resolve_client_id(code, env=env)
    if cid is None:
        return result
    vm = msdmod.load_vendor_master(cid, env=env)
    result["in_msd"] = True
    result["msd_vendors"] = len(vm)
    result["coherence"] = msdmod.coherence_summary(vm)
    result["vendor_master"] = vm
    return result


# ------------------------------------------------------------------------------ per client ----
def profile_client(client_key, env):
    cfg = clientcfg.load_config(client_key)
    name = clientcfg.display_name(cfg, client_key)
    out = {"client_key": client_key, "client": name, "status": "", "notes": [],
           "tables": None, "columns": None, "taxonomy_levels": None,
           "taxonomy_detail": None, "top_desc": None, "msd": None,
           "desc_fields": None, "desc_field_detail": None}

    print(f"\n== {name} ({client_key}) " + "=" * (50 - len(name)))
    try:
        cn = connect_client(client_key, env=env)
    except Exception as e:
        out["status"] = "CONNECT FAILED"
        out["notes"].append(str(e).strip().splitlines()[0][:300])
        print(f"  connect: FAILED - {out['notes'][-1]}")
        return out
    print("  connect: OK")

    try:
        table = clientcfg.source_table(cfg)
        cols = clientcfg.column_map(cfg)

        if not table:
            out["status"] = "NO TABLE CONFIGURED"
            out["notes"].append("source.table not set - listing candidate tables/views instead")
            out["tables"] = discover_tables(cn)
            print(f"  no source.table - discovered {len(out['tables'])} candidate objects")
            return out

        out["columns"] = table_columns(cn, table)

        # One check, from the loader, so profiling and a real run can never disagree about
        # what "configured" means.
        problems = clientcfg.missing_for_run(cfg)
        if problems:
            out["status"] = "CONFIG INCOMPLETE"
            out["notes"].extend(problems)
            print("  config incomplete: " + "; ".join(problems))
            return out

        fields = clientcfg.description_fields(cfg)
        d = desc_expr(fields)          # the composed description - what the judge reads
        out["desc_fields"] = fields

        vol = profile_volume(cn, table, cols, d)
        out.update({k: vol[k] for k in
                    ["lines", "spend", "spend_abs", "vendors",
                     "spend_unparseable", "spend_negative", "credit_spend_abs", "cat0_null"]})
        print(f"  lines={vol['lines']:,}  spend=${(vol['spend'] or 0):,.0f}  "
              f"vendors={vol['vendors']:,}")
        print(f"  credits: {int(vol['spend_negative']):,} negative lines, "
              f"${(vol['credit_spend_abs'] or 0):,.0f} gross")
        if vol["spend_unparseable"]:
            print(f"  ** spend unparseable on {int(vol['spend_unparseable']):,} lines")

        units, units_vt = profile_units(cn, table, cols, d)
        out["judge_units"] = units
        out["units_vendor_term"] = units_vt
        out["case_b_units"] = units - units_vt
        out["dup_factor"] = (vol["lines"] / units) if units else 0
        print(f"  judge units={units:,} (dedup {out['dup_factor']:.1f}x)  "
              f"same-term-two-categories extra units={out['case_b_units']:,}")

        n = max(int(vol["lines"]), 1)
        out["desc_blank_pct"] = 100.0 * int(vol["blank"]) / n
        out["desc_short_pct"] = 100.0 * int(vol["very_short"]) / n
        out["desc_numeric_pct"] = 100.0 * int(vol["numeric_only"]) / n
        out["pii_digit_run_pct"] = 100.0 * int(vol["long_digit_run"]) / n
        out["pii_marker_lines"] = int(vol["pii_marker"])
        print(f"  descriptions: blank={out['desc_blank_pct']:.1f}%  "
              f"short={out['desc_short_pct']:.1f}%  numeric-only={out['desc_numeric_pct']:.1f}%")
        if out["pii_marker_lines"]:
            print(f"  ** PII markers (MRN/UR/DOB) on {out['pii_marker_lines']:,} lines")

        out["desc_field_detail"] = profile_desc_fields(cn, table, fields)
        if out["desc_field_detail"] is not None and not out["desc_field_detail"].empty:
            for r2 in out["desc_field_detail"].itertuples(index=False):
                print(f"    {r2.field[:44]:46} populated {r2.populated_pct:5.1f}%  "
                      f"supplied {r2.supplied_the_text:,}")

        out["taxonomy_levels"] = taxonomy_by_level(cn, table, cols)
        out["taxonomy_detail"] = taxonomy_detail(cn, table, cols)
        out["top_desc"] = top_descriptions(cn, table, cols, d)

        # generic-description share, measured over the top descriptions by spend
        td = out["top_desc"]
        if td is not None and not td.empty:
            gen = td[td["item_desc"].fillna("").str.strip().isin(GENERIC_TERMS)]
            tot = td["spend"].sum()
            out["generic_spend_pct"] = 100.0 * gen["spend"].sum() / tot if tot else 0.0
        out["status"] = "OK"

    except Exception as e:
        out["status"] = "PROFILE FAILED"
        out["notes"].append(f"{type(e).__name__}: {e}")
        traceback.print_exc()
    finally:
        cn.close()

    try:
        out["msd"] = profile_msd(cfg, env)
        m = out["msd"]
        print(f"  MSD: code={m['msd_client_code']}  in_msd={m['in_msd']}  vendors={m['msd_vendors']:,}")
    except Exception as e:
        out["notes"].append(f"MSD check failed: {type(e).__name__}: {e}")
        print(f"  MSD: check failed - {e}")

    return out


# -------------------------------------------------------------------------------- workbook ----
def write_workbook(results):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    NAVY, WHITE = "1F3864", "FFFFFF"
    wb = Workbook()

    def header(ws, row, headers, widths):
        for j, (h, w) in enumerate(zip(headers, widths), 1):
            c = ws.cell(row, j, h)
            c.font = Font(bold=True, color=WHITE)
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(j)].width = w

    def sheet(title):
        ws = wb.create_sheet(title[:31])
        ws.sheet_view.showGridLines = False
        return ws

    def dump(ws, df, start_row, widths=None):
        if df is None or df.empty:
            ws.cell(start_row, 1, "(no data)")
            return start_row + 2
        widths = widths or [max(12, min(45, len(str(c)) + 6)) for c in df.columns]
        header(ws, start_row, list(df.columns), widths)
        for i, row in enumerate(df.itertuples(index=False), start_row + 1):
            for j, v in enumerate(row, 1):
                cell = ws.cell(i, j, v)
                if isinstance(v, float) and df.columns[j - 1] in ("spend", "spend_total"):
                    cell.number_format = "$#,##0"
        return start_row + len(df) + 3

    # ---- Summary -------------------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Medical QA (Indirects) - Phase 0 client profiling"
    ws["A1"].font = Font(size=14, bold=True, color=NAVY)
    ws["A2"] = f"Generated {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = Font(size=9, color="808080")

    cols = ["client", "status", "lines", "spend", "vendors", "judge_units", "dup_factor",
            "case_b_units", "credit_lines", "credit_spend_abs", "cat0_null",
            "desc_blank_pct", "desc_short_pct", "desc_numeric_pct",
            "generic_spend_pct", "pii_marker_lines", "msd_in", "msd_vendors", "notes"]
    rows = []
    for r in results:
        m = r.get("msd") or {}
        rows.append({
            "client": r["client"], "status": r["status"],
            "lines": r.get("lines"), "spend": r.get("spend"), "vendors": r.get("vendors"),
            "judge_units": r.get("judge_units"), "dup_factor": r.get("dup_factor"),
            "case_b_units": r.get("case_b_units"),
            "credit_lines": r.get("spend_negative"),
            "credit_spend_abs": r.get("credit_spend_abs"),
            "cat0_null": r.get("cat0_null"),
            "desc_blank_pct": r.get("desc_blank_pct"), "desc_short_pct": r.get("desc_short_pct"),
            "desc_numeric_pct": r.get("desc_numeric_pct"),
            "generic_spend_pct": r.get("generic_spend_pct"),
            "pii_marker_lines": r.get("pii_marker_lines"),
            "msd_in": m.get("in_msd"), "msd_vendors": m.get("msd_vendors"),
            "notes": "; ".join(r.get("notes") or []),
        })
    summary = pd.DataFrame(rows, columns=cols)
    dump(ws, summary, 4,
         widths=[26, 20, 12, 16, 11, 13, 11, 13, 12, 16, 12,
                 12, 12, 13, 13, 13, 9, 12, 60])

    ws.cell(len(summary) + 8, 1,
            "POC selection: prefer the client with the LOWEST blank/short/numeric-only "
            "description share - the judge reads descriptions, so poor text caps achievable "
            "accuracy before any modelling. Then prefer moderate line volume and the highest "
            "MSD coverage.").font = Font(italic=True, color="808080")

    # ---- per client ------------------------------------------------------------------------
    for r in results:
        ws = sheet(r["client_key"])
        ws["A1"] = f"{r['client']} - {r['status']}"
        ws["A1"].font = Font(size=13, bold=True, color=NAVY)
        row = 3
        if r.get("notes"):
            for nte in r["notes"]:
                ws.cell(row, 1, nte).font = Font(color="A13030")
                row += 1
            row += 1
        if r.get("tables") is not None:
            ws.cell(row, 1, "Candidate tables / views (source.table not set)").font = Font(bold=True)
            row = dump(ws, r["tables"], row + 1)
        if r.get("columns") is not None:
            ws.cell(row, 1, "Columns in the configured table").font = Font(bold=True)
            row = dump(ws, r["columns"], row + 1)
        if r.get("desc_field_detail") is not None and not r["desc_field_detail"].empty:
            ws.cell(row, 1, "Description fields - fill rate and which one wins").font = Font(bold=True)
            row += 1
            ws.cell(row, 1, "Priority order is set in config; first non-blank wins. "
                            "'supplied the text' is how often each field actually won.")\
                .font = Font(italic=True, color="808080")
            row = dump(ws, r["desc_field_detail"], row + 1, widths=[46, 16, 14, 18])
        if r.get("taxonomy_levels") is not None:
            ws.cell(row, 1, "Taxonomy shape by level").font = Font(bold=True)
            row = dump(ws, r["taxonomy_levels"], row + 1)
        m = r.get("msd") or {}
        if isinstance(m.get("coherence"), pd.DataFrame) and not m["coherence"].empty:
            ws.cell(row, 1, "MSD coherence split").font = Font(bold=True)
            row = dump(ws, m["coherence"], row + 1)

        if r.get("taxonomy_detail") is not None and not r["taxonomy_detail"].empty:
            ws2 = sheet(f"{r['client_key'][:22]}-taxonomy")
            ws2["A1"] = f"{r['client']} - every category by spend"
            ws2["A1"].font = Font(size=12, bold=True, color=NAVY)
            ws2["A2"] = ("Read this to choose clinical_gate.exclude_categories. "
                         "The clinical gate is CATEGORY-BASED ONLY - never keywords.")
            ws2["A2"].font = Font(italic=True, color="808080")
            dump(ws2, r["taxonomy_detail"], 4, widths=[10, 55, 12, 18])

        if r.get("top_desc") is not None and not r["top_desc"].empty:
            ws3 = sheet(f"{r['client_key'][:20]}-topdesc")
            ws3["A1"] = f"{r['client']} - top descriptions by spend"
            ws3["A1"].font = Font(size=12, bold=True, color=NAVY)
            ws3["A2"] = "This is what the judge will be reading. If it's unreadable here, it's unjudgeable."
            ws3["A2"].font = Font(italic=True, color="808080")
            dump(ws3, r["top_desc"], 4, widths=[70, 12, 18])

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    wb.save(OUT_PATH)
    return OUT_PATH


def main():
    env = load_env()
    wanted = sys.argv[1:] or clientcfg.list_clients()
    if not wanted:
        print("No clients found under clients/.")
        return

    print(f"Profiling {len(wanted)} client(s): {', '.join(wanted)}")
    try:
        cn = connect_msd(env=env)
        cn.close()
        print("MSD connection: OK")
    except Exception as e:
        print(f"MSD connection: FAILED - {e}")

    results = [profile_client(k, env) for k in wanted]
    out = write_workbook(results)
    print(f"\nWrote {out}")

    ok = [r for r in results if r["status"] == "OK"]
    if ok:
        ranked = sorted(ok, key=lambda r: (r.get("desc_blank_pct", 100)
                                           + r.get("desc_short_pct", 100)
                                           + r.get("desc_numeric_pct", 100)))
        print("\nPOC candidates by description quality (best first):")
        for r in ranked:
            print(f"  {r['client']:28} unusable-desc "
                  f"{r.get('desc_blank_pct', 0) + r.get('desc_short_pct', 0) + r.get('desc_numeric_pct', 0):5.1f}%  "
                  f"units={r.get('judge_units', 0):,}")


if __name__ == "__main__":
    main()
