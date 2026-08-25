"""Flat Excel copy of the browsable taxonomy chart - every category, all five levels, repeats kept.

    python pipeline/taxonomy_draft.py

Reads the newest MERGED workbook, opens no database, writes
`output/Taxonomy/Taxonomy_Draft (Indirect).xlsx`.

WHY IT EXISTS. Sameer, 2026-08-13: *"create a new excel file which is nothing but a copy of the
html with all levels, doesnt matter if its repeated, call it Taxonomy_Draft"*. The HIERARCHY chart
is the thing he reads the taxonomy in; this is the same content in the form he can sort, filter and
mark up.

IT IS A COPY, AND NOTHING MAY EXIST ONLY IN IT. The DECISIONS workbook is the one place a ruling
lives - a second workbook holding the same tree is a second place an answer could be written, which
is exactly the failure the four-file rule was written to prevent. Regenerated from the MERGED file
on demand and overwritten without ceremony; if a ruling is ever typed into it, it stops being
regenerable and has to be treated like a returned review workbook.

The level columns repeat the deepest label down to Level 4 - the padding convention from PLAN.md
change 175 - because that is how the source hospitals pad and how the chart displays.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from merge_taxonomy import OUTDIR   # noqa: E402
from taxonomy_audit import newest_merged   # noqa: E402

NAME = "Taxonomy_Draft (Indirect).xlsx"
COLS = ["LEVEL_0", "LEVEL_1", "LEVEL_2", "LEVEL_3", "LEVEL_4", "LINES"]


def main():
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    src = newest_merged()
    ws_in = load_workbook(src, data_only=True).active
    hdr = [str(c.value or "").strip() for c in ws_in[1]]
    idx = [hdr.index(c) for c in COLS]

    rows = [[r[i] for i in idx] for r in ws_in.iter_rows(min_row=2, values_only=True) if r[0]]
    rows.sort(key=lambda r: [str(x or "") for x in r[:5]])

    wb = Workbook()
    ws = wb.active
    ws.title = "Taxonomy"
    ws.append(COLS)
    for r in rows:
        ws.append(r)

    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="44546A")
    for c in ws[1]:
        c.font, c.fill = head, fill
        c.alignment = Alignment(vertical="center")
    for col, width in zip("ABCDEF", (16, 30, 34, 38, 38, 11)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(rows) + 1}"
    for c in ws["F"][1:]:
        c.number_format = "#,##0"

    out = os.path.join(OUTDIR, NAME)
    wb.save(out)
    print(f"  read   : {os.path.basename(src)}")
    print(f"  WRITTEN: {out}   ({len(rows)} categories, {sum(r[5] or 0 for r in rows):,} lines)")


if __name__ == "__main__":
    main()
