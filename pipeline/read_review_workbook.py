"""Read a COMPLETED review workbook back into the review layer of `qa_line`, and report the
agreement rate - the figure that gates publishing any accuracy number.

WHAT THIS WRITES, AND WHAT IT MUST NEVER WRITE
----------------------------------------------
Writes ONLY: REVIEW_STATUS, REVIEWED_BY, REVIEWED_AT, REVIEW_OVERRIDE_VERDICT,
REVIEW_OVERRIDE_CATEGORY, REVIEW_NOTE.

Never writes VERDICT, CONFIDENCE, BASIS, RATIONALE or any SUGGESTED_CATEGORY_* column. The finding
is immutable - the analyst redirects, they never rewrite - and keeping both layers side by side is
the only reason the override rate measures anything. On 2026-08-05 nine verdicts were overwritten
by a batch that assumed what it was holding, and the guard written to catch it could not fail. So
this script does not rely on care: it CHECKSUMS the finding columns before and after the write and
refuses to commit if they moved.

Dry run by DEFAULT. Nothing is written without --commit.

Usage:
    python pipeline/read_review_workbook.py --file "<completed workbook>.xlsx"
    python pipeline/read_review_workbook.py --file "<...>.xlsx" --reviewer "Monali" --commit
"""
import argparse
import datetime
import os

from openpyxl import load_workbook

import clientcfg
from db import connect_qa, load_env
from make_review_workbook import COLUMNS, RESPONSE_OPTIONS

AGREE, DISAGREE_CURRENT_OK, DISAGREE_NEW_CAT = RESPONSE_OPTIONS

STATUS = {
    AGREE: "agreed",
    DISAGREE_CURRENT_OK: "disagreed_current_ok",
    DISAGREE_NEW_CAT: "disagreed_new_category",
}

COL_ID = 1
COL_RESPONSE = len(COLUMNS) - 2
COL_CATEGORY = len(COLUMNS) - 1
COL_NOTE = len(COLUMNS)
COL_VERDICT = [i for i, (_, s, _) in enumerate(COLUMNS, start=1) if s == "VERDICT"][0]

FINDING_COLS = ["VERDICT", "CONFIDENCE", "BASIS", "RATIONALE",
                "SUGGESTED_CATEGORY_LVL_0", "SUGGESTED_CATEGORY_LVL_1",
                "SUGGESTED_CATEGORY_LVL_2", "SUGGESTED_CATEGORY_LVL_3",
                "SUGGESTED_CATEGORY_LVL_4"]


def finding_checksum(cur, client):
    """A fingerprint of OUR layer. It must be identical before and after the write."""
    cols = " + '|' + ".join(f"ISNULL(CAST([{c}] AS nvarchar(4000)),'~')" for c in FINDING_COLS)
    cur.execute(f"""SELECT CHECKSUM_AGG(BINARY_CHECKSUM({cols})), COUNT(*)
                      FROM qa_line WHERE client_code = ?""", client)
    return cur.fetchone()


def read_sheet(path):
    wb = load_workbook(path, data_only=True)
    if "Review" not in wb.sheetnames:
        raise SystemExit(f"STOP: no 'Review' sheet in {os.path.basename(path)}.")
    rv = wb["Review"]

    expected = [h for h, _, _ in COLUMNS]
    actual = [rv.cell(row=1, column=c).value for c in range(1, len(expected) + 1)]
    if actual != expected:
        raise SystemExit(
            "STOP: the workbook's columns are not the ones it was generated with - a column has "
            "been inserted, deleted or renamed, so answers cannot be trusted to line up.\n"
            f"  expected: {expected}\n  found   : {actual}")

    rows = []
    for r in range(2, rv.max_row + 1):
        line_id = rv.cell(row=r, column=COL_ID).value
        if line_id is None:
            continue
        rows.append({
            "excel_row": r,
            "line_id": int(line_id),
            "response": (rv.cell(row=r, column=COL_RESPONSE).value or "").strip(),
            "category": (rv.cell(row=r, column=COL_CATEGORY).value or "").strip(),
            "note": (rv.cell(row=r, column=COL_NOTE).value or "").strip() or None,
        })
    return rows


def main(path, reviewer, commit):
    load_env()
    rows = read_sheet(path)
    qa = connect_qa()
    cur = qa.cursor()

    ids = [r["line_id"] for r in rows]
    marks = ",".join("?" * len(ids))
    cur.execute(f"""SELECT qa_line_id, client_code, VERDICT,
                           CATEGORY_LVL_0, CATEGORY_LVL_1, CATEGORY_LVL_2,
                           CATEGORY_LVL_3, CATEGORY_LVL_4
                      FROM qa_line WHERE qa_line_id IN ({marks})""", *ids)
    db = {}
    for r in cur.fetchall():
        db[r[0]] = {
            "client": r[1], "verdict": r[2],
            "current_path": " > ".join((x or "").strip() for x in r[3:8]),
            "uncategorised": (r[4] or "").strip() == clientcfg.LEVEL_UNCATEGORISED,
        }

    missing = [r["line_id"] for r in rows if r["line_id"] not in db]
    if missing:
        raise SystemExit(f"STOP: {len(missing)} row(s) carry a line id that is not in qa_line, "
                         f"e.g. {missing[:5]}. Was this workbook built from this run?")

    clients = {v["client"] for v in db.values()}
    if len(clients) != 1:
        raise SystemExit(f"STOP: the file mixes clients: {clients}. One hospital per workbook.")
    client = clients.pop()

    cur.execute("""SELECT path_full, in_scope FROM qa_category WHERE client_code = ?""", client)
    tax = {}
    for p, s in cur.fetchall():
        if p:
            tax[p.strip()] = tax.get(p.strip(), False) or bool(s)

    # ---------------------------------------------------------------- validate, then decide
    updates, problems, blank = [], [], 0
    for r in rows:
        d = db[r["line_id"]]
        resp, cat = r["response"], r["category"]

        if not resp:
            if cat or r["note"]:
                problems.append((r["excel_row"], r["line_id"],
                                 "a category or note was entered but no response was chosen"))
            else:
                blank += 1
            continue
        if resp not in STATUS:
            problems.append((r["excel_row"], r["line_id"], f"response not recognised: {resp!r}"))
            continue

        if resp == AGREE and cat:
            problems.append((r["excel_row"], r["line_id"],
                             "marked Agree but a category was also picked - which is meant?"))
            continue

        if resp == DISAGREE_NEW_CAT and not cat:
            # Excel cannot force this (one validation rule per cell); the app will. Report it.
            problems.append((r["excel_row"], r["line_id"],
                             "disagreed and said a category would follow, but none was picked"))
            continue

        if resp == DISAGREE_CURRENT_OK and d["verdict"] == "Correct":
            problems.append((r["excel_row"], r["line_id"],
                             "'current category is fine' on a line we already called Correct - "
                             "that is agreement, not disagreement"))
            continue

        if resp == DISAGREE_CURRENT_OK and d["uncategorised"]:
            problems.append((r["excel_row"], r["line_id"],
                             "'current category is fine' on a line that has no category at all"))
            continue

        if cat:
            if cat not in tax:
                problems.append((r["excel_row"], r["line_id"],
                                 f"category is not in {client}'s taxonomy: {cat[:70]!r}"))
                continue
            # the per-line scope guard, same rule the judge applies: only an uncategorised line
            # may be sent outside the indirect branches
            if not tax[cat] and not d["uncategorised"]:
                problems.append((r["excel_row"], r["line_id"],
                                 "category is outside the indirect branches and this line "
                                 "already has a category of its own"))
                continue

        if resp == AGREE:
            override_verdict, override_cat = None, None
        elif resp == DISAGREE_CURRENT_OK:
            override_verdict, override_cat = "Correct", d["current_path"]
        else:
            override_verdict, override_cat = "Incorrect", cat

        updates.append((STATUS[resp], reviewer, override_verdict, override_cat,
                        r["note"], r["line_id"]))

    # ---------------------------------------------------------------- report
    print(f"\n  file      : {os.path.basename(path)}")
    print(f"  client    : {client}")
    print(f"  reviewer  : {reviewer}")
    print(f"  rows      : {len(rows):,}   answered: {len(updates):,}   "
          f"left blank: {blank:,}   problems: {len(problems):,}")

    if problems:
        print(f"\n  PROBLEMS - these rows are NOT written, and are reported rather than guessed at:")
        for xr, lid, why in problems[:25]:
            print(f"     Excel row {xr:>4} (line {lid}): {why}")
        if len(problems) > 25:
            print(f"     ... and {len(problems) - 25} more")

    if updates:
        by_status = {}
        for u in updates:
            by_status[u[0]] = by_status.get(u[0], 0) + 1
        agreed = by_status.get("agreed", 0)
        disagreed = len(updates) - agreed
        print(f"\n  AGREEMENT: {agreed:,} of {len(updates):,} answered "
              f"({100.0 * agreed / len(updates):.1f}%)")
        for k, v in sorted(by_status.items()):
            print(f"     {k:<26} {v:>5}")

        print(f"\n  agreement by OUR verdict - where the judge is trusted, and where it is not:")
        per = {}
        for u in updates:
            v = db[u[5]]["verdict"]
            a, t = per.get(v, (0, 0))
            per[v] = (a + (1 if u[0] == "agreed" else 0), t + 1)
        for v in ("Incorrect", "Uncertain", "Correct"):
            if v in per:
                a, t = per[v]
                print(f"     we said {v:<10} agreed {a:>4} of {t:>4}  ({100.0*a/t:.1f}%)")
        if "Correct" in per:
            a, t = per["Correct"]
            if t - a:
                print(f"     -> {t - a} disagreement(s) on lines we called Correct: judge "
                      f"false positives, the most valuable finding in the file")

    if not commit:
        print("\n  DRY RUN - nothing written. Re-run with --commit to write these answers.")
        qa.close()
        return

    if not updates:
        print("\n  nothing to write.")
        qa.close()
        return

    before = finding_checksum(cur, client)
    cur.executemany("""UPDATE qa_line
                          SET REVIEW_STATUS = ?, REVIEWED_BY = ?, REVIEWED_AT = SYSUTCDATETIME(),
                              REVIEW_OVERRIDE_VERDICT = ?, REVIEW_OVERRIDE_CATEGORY = ?,
                              REVIEW_NOTE = ?
                        WHERE qa_line_id = ?""", updates)
    after = finding_checksum(cur, client)

    if before != after:
        qa.rollback()
        raise SystemExit(
            f"STOP - ROLLED BACK. The finding columns changed during the write "
            f"(before={before}, after={after}). Nothing has been committed. This should be "
            f"impossible; investigate before retrying.")

    qa.commit()
    print(f"\n  COMMITTED {len(updates):,} review answers for {client}.")
    print(f"  finding-column checksum unchanged before and after: {before[0]} "
          f"over {before[1]:,} rows - our verdicts, rationales and suggestions were not touched.")
    cur.execute("""SELECT COUNT(*) FROM qa_line
                    WHERE client_code = ? AND REVIEW_STATUS IS NOT NULL""", client)
    print(f"  rows now carrying a review answer for this client: {cur.fetchone()[0]:,}")
    qa.close()


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Ingest a completed review workbook.")
    ap.add_argument("--file", required=True, help="path to the completed .xlsx")
    ap.add_argument("--reviewer", default=None,
                    help="who reviewed it - stored in REVIEWED_BY (required with --commit)")
    ap.add_argument("--commit", action="store_true", help="write; without this it is a dry run")
    a = ap.parse_args()
    if a.commit and not a.reviewer:
        raise SystemExit("STOP: --reviewer is required with --commit. A review with no author "
                         "cannot be audited, and REVIEWED_BY is what makes the override rate "
                         "attributable.")
    main(a.file, a.reviewer or "(dry run)", a.commit)
